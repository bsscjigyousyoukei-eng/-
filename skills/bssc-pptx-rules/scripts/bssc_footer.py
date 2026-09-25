"""BSSC: 右下のフッター（自動連動ページ番号＋著作権表示）を付与する。
何度実行しても同じ結果になる（冪等）。2-3. のフォント是正処理のあとに実行する。

使い方（著作権表示の有無は、0. の確認に対するユーザーの回答どおりに必ず指定する）:
    python3 bssc_footer.py output.pptx --copyright      # 掲載する
    python3 bssc_footer.py output.pptx --no-copyright   # 掲載しない
"""
import argparse, uuid
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn

PN_NAME, CR_NAME = "BSSC_PageNumber", "BSSC_Copyright"
CR_TEXT = "© BSSC Co., Ltd."
Y, H = Inches(6.95), Inches(0.35)
PN_X, PN_W = Inches(10.10), Inches(0.55)          # 最も右：ページ番号
CR_X, CR_W = Inches(7.45), Inches(2.50)           # その左：著作権表示（間隔 0.15in）
NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')

def _sp(shape_id, name, x, w, inner):
    return etree.fromstring(f'''<p:sp {NS}>
  <p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="{x}" y="{Y}"/><a:ext cx="{w}" cy="{H}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>
  <p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="b"/><a:lstStyle/>
    <a:p><a:pPr algn="r"/>{inner}</a:p></p:txBody></p:sp>''')

def _rpr(font, sz):
    return (f'<a:rPr lang="ja-JP" sz="{sz}" b="1"><a:latin typeface="{font}"/>'
            f'<a:ea typeface="{font}"/></a:rPr>')

def _page_number(shape_id):
    guid = "{" + str(uuid.uuid4()).upper() + "}"
    return _sp(shape_id, PN_NAME, PN_X, PN_W,
               f'<a:fld id="{guid}" type="slidenum">{_rpr("MS Pゴシック", 1600)}<a:t>‹#›</a:t></a:fld>')

def _copyright(shape_id):
    return _sp(shape_id, CR_NAME, CR_X, CR_W, f'<a:r>{_rpr("メイリオ", 1100)}<a:t>{CR_TEXT}</a:t></a:r>')

def _next_id(tree):
    ids = [int(e.get("id")) for e in tree.iter(qn("p:cNvPr")) if e.get("id", "").isdigit()]
    return max(ids + [1]) + 1

def _add(tree, factory):
    ext = tree.find(qn("p:extLst"))
    el = factory(_next_id(tree))
    ext.addprevious(el) if ext is not None else tree.append(el)

def _remove(shapes, pred):
    for shp in [x for x in shapes if pred(x)]:
        shp._element.getparent().remove(shp._element)

def _is_old_footer(sh, sw, sh_h):
    """スライド上の旧フッター（固定数字のページ番号・sldNumプレースホルダ・直書きの著作権表示・前回付与分）"""
    if sh.name in (PN_NAME, CR_NAME):
        return True
    if sh.is_placeholder and sh.placeholder_format.type is not None \
            and sh.placeholder_format.type.name == "SLIDE_NUMBER":
        return True
    if sh.has_text_frame:
        t = sh.text_frame.text.strip()
        if t.replace(" ", "") == CR_TEXT.replace(" ", ""):
            return True
        if t.isdigit() and sh.left is not None and sh.left > sw * 0.8 and sh.top > sh_h * 0.85:
            return True
    return False

def apply(path, copyright_on):
    prs = Presentation(path)
    sw, sh_h = prs.slide_width, prs.slide_height
    # 1) 表紙を0番とし、2枚目が「1」になるよう開始番号を0にする
    prs.part._element.set("firstSlideNum", "0")
    # 2) 全レイアウトにページ番号（と著作権表示）を1つずつ置き直す → 追加スライドにも自動で出る
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            tree = layout.shapes._spTree
            _remove(layout.shapes, lambda x: x.name in (PN_NAME, CR_NAME))
            _add(tree, _page_number)
            if copyright_on:
                _add(tree, _copyright)
    # 3) スライド上の旧フッターを削除（二重表示防止）
    for s in prs.slides:
        _remove(s.shapes, lambda x: _is_old_footer(x, sw, sh_h))
    # 4) 表紙：レイアウトの図形を非表示にして番号を消し、著作権表示だけ表紙に直接置く
    cover = prs.slides[0]
    cover._element.set("showMasterSp", "0")
    if copyright_on:
        _add(cover.shapes._spTree, _copyright)
    extra = [x.name for x in cover.slide_layout.shapes
             if not x.is_placeholder and x.name not in (PN_NAME, CR_NAME)]
    if extra:
        print("注意: 表紙レイアウトの装飾図形も非表示になります:", extra)
    prs.save(path)

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--copyright", dest="cr", action="store_true")
    g.add_argument("--no-copyright", dest="cr", action="store_false")
    a = ap.parse_args()
    apply(a.path, a.cr)
