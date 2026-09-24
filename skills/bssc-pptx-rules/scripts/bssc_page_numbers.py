"""BSSC: 自動連動ページ番号（スライド番号フィールド方式）を付与する。
何度実行しても同じ結果になる（冪等）。フォント是正処理のあとに実行する。"""
import sys, uuid
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.ns import qn

NAME = "BSSC_PageNumber"
X, Y, W, H = Inches(10.10), Inches(6.95), Inches(0.55), Inches(0.35)
NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')

def _box(shape_id):
    return etree.fromstring(f'''<p:sp {NS}>
  <p:nvSpPr><p:cNvPr id="{shape_id}" name="{NAME}"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>
  <p:spPr><a:xfrm><a:off x="{X}" y="{Y}"/><a:ext cx="{W}" cy="{H}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>
  <p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="b"/><a:lstStyle/>
    <a:p><a:pPr algn="r"/><a:fld id="{{{str(uuid.uuid4()).upper()}}}" type="slidenum">
      <a:rPr lang="ja-JP" sz="1600" b="1"><a:latin typeface="MS Pゴシック"/><a:ea typeface="MS Pゴシック"/></a:rPr>
      <a:t>‹#›</a:t></a:fld></a:p></p:txBody></p:sp>''')

def _next_id(spTree):
    ids = [int(e.get("id")) for e in spTree.iter(qn("p:cNvPr")) if e.get("id", "").isdigit()]
    return max(ids + [1]) + 1

def _is_old_number(sh, sw, sh_h):
    """スライド上の旧ページ番号（固定数字の箱・sldNumプレースホルダ・前回付与分）を判定"""
    if sh.name == NAME:
        return True
    if sh.is_placeholder and sh.placeholder_format.type is not None \
            and sh.placeholder_format.type.name == "SLIDE_NUMBER":
        return True
    if sh.has_text_frame and sh.text_frame.text.strip().isdigit() \
            and sh.left is not None and sh.left > sw * 0.8 and sh.top > sh_h * 0.85:
        return True
    return False

def apply(path):
    prs = Presentation(path)
    sw, sh_h = prs.slide_width, prs.slide_height
    # 1) 表紙を0番とし、2枚目が「1」になるよう開始番号を0にする
    prs.part._element.set("firstSlideNum", "0")
    # 2) 全レイアウトに番号フィールドを1つずつ置く（新規追加スライドにも自動で出る）
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            tree = layout.shapes._spTree
            for e in [s._element for s in layout.shapes if s.name == NAME]:
                tree.remove(e)
            ext = tree.find(qn("p:extLst"))
            box = _box(_next_id(tree))
            ext.addprevious(box) if ext is not None else tree.append(box)
    # 3) スライド上の旧ページ番号を削除（二重表示防止）
    for s in prs.slides:
        for shp in [x for x in s.shapes if _is_old_number(x, sw, sh_h)]:
            shp._element.getparent().remove(shp._element)
    # 4) 表紙はマスター/レイアウトの図形を非表示にして番号を出さない
    cover = prs.slides[0]
    cover._element.set("showMasterSp", "0")
    extra = [x.name for x in cover.slide_layout.shapes
             if not x.is_placeholder and x.name != NAME]
    if extra:
        print("注意: 表紙レイアウトの装飾図形も非表示になります:", extra)
    prs.save(path)

if __name__ == "__main__":
    apply(sys.argv[1])
